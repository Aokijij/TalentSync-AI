import csv
import hashlib
import io
import re
import unicodedata
from datetime import datetime, timedelta
from urllib.parse import urlparse

from openpyxl import load_workbook

from app.application.errors import CatalogProviderError, UseCaseError
from app.application.ports.services import JobCatalog, TextAnalysis
from app.application.ports.unit_of_work import UnitOfWork
from app.domain.entities.catalog import CatalogJob
from app.domain.entities.job_sectors import JOB_SECTORS
from app.domain.entities.records import User

MAX_IMPORT_ROWS = 2_000
IMPORT_HEADERS = (
    "fuente",
    "id_externo",
    "empresa",
    "cargo",
    "descripcion",
    "requisitos",
    "sector",
    "enlace_externo",
    "fecha_publicacion",
    "fecha_vencimiento",
    "salario",
    "departamento",
    "ciudad",
    "modalidad",
    "tipo_contrato",
    "habilidades",
    "beneficios",
    "idiomas",
    "sitio_empresa",
    "descripcion_empresa",
)


def _key(value: object) -> str:
    text = "".join(
        character
        for character in unicodedata.normalize("NFKD", str(value or ""))
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _text(row: dict[str, object], name: str, *, required: bool = False) -> str:
    value = " ".join(str(row.get(name) or "").strip().split())
    if required and not value:
        raise ValueError(f"Falta la columna obligatoria “{name}”")
    return value


def _split(value: object) -> list[str]:
    return list(
        dict.fromkeys(
            item.strip().lower()
            for item in re.split(r"[|,]", str(value or ""))
            if item.strip()
        )
    )


def _salary(value: object) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().lower()
    if any(term in raw for term in ("a convenir", "por definir", "no especific")):
        return None
    amounts: list[float] = []
    for token in re.findall(r"\d+(?:[.,]\d+)*", raw):
        groups = re.split(r"[.,]", token)
        if len(groups) > 1 and all(len(group) == 3 for group in groups[1:]):
            amounts.append(float("".join(groups)))
            continue
        normalized = token.replace(",", ".")
        try:
            amount = float(normalized)
        except ValueError:
            continue
        if "millon" in raw and amount < 1_000:
            amount *= 1_000_000
        amounts.append(amount)
    if not amounts:
        raise ValueError("El salario debe contener un valor numérico")
    # The current model stores one value. For a range, keep its lower bound so
    # salary filters do not promise more than the source actually offers.
    return min(amounts)


def _date(value: object, field: str) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    raw = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(raw, pattern)
        except ValueError:
            continue
    raise ValueError(f"{field} debe usar AAAA-MM-DD")


def _url(value: object, field: str, *, required: bool = False) -> str | None:
    raw = str(value or "").strip()
    if not raw:
        if required:
            raise ValueError(f"Falta la columna obligatoria “{field}”")
        return None
    parsed = urlparse(raw)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError(f"{field} debe ser una dirección web válida")
    if len(raw) > 1_000:
        raise ValueError(f"{field} es demasiado largo")
    return raw


def _languages(value: object) -> list[dict[str, str]]:
    languages: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in re.split(r"[|,]", str(value or "")):
        if not item.strip():
            continue
        name, separator, level = item.partition(":")
        name = " ".join(name.lower().strip().split())
        level = level.strip().upper()
        if not separator or level not in {"A1", "A2", "B1", "B2", "C1", "C2", "NATIVE"}:
            raise ValueError(
                "Los idiomas deben usar el formato idioma:nivel, por ejemplo inglés:B2"
            )
        if name in seen:
            raise ValueError(f"El idioma “{name}” está repetido")
        seen.add(name)
        languages.append({"name": name, "level": level})
    return languages


def _rows(filename: str, content: bytes) -> list[dict[str, object]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise UseCaseError(
                status_code=400, detail="El CSV debe estar guardado en formato UTF-8"
            ) from error
        try:
            dialect = csv.Sniffer().sniff(decoded[:4_096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
        if not reader.fieldnames:
            raise UseCaseError(
                status_code=400, detail="El archivo no contiene encabezados"
            )
        return [
            {_key(name): value for name, value in row.items() if name}
            for row in reader
            if any(str(value or "").strip() for value in row.values())
        ]
    if suffix == "xlsx":
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [_key(value) for value in next(values, ())]
            if not any(headers):
                raise ValueError("El archivo no contiene encabezados")
            return [
                {
                    headers[index]: value
                    for index, value in enumerate(row)
                    if index < len(headers)
                }
                for row in values
                if any(str(value or "").strip() for value in row)
            ]
        except (OSError, ValueError, TypeError) as error:
            raise UseCaseError(
                status_code=400, detail=f"El Excel no es válido: {error}"
            ) from error
    raise UseCaseError(status_code=400, detail="Selecciona un archivo CSV o XLSX")


def _mapped(value: str, mapping: dict[str, str], default: str, field: str) -> str:
    if not value:
        return default
    result = mapping.get(_key(value))
    if result is None:
        raise ValueError(f"Valor no reconocido en {field}: {value}")
    return result


def _sector(value: str) -> str:
    sectors = {_key(item): item for item in JOB_SECTORS}
    result = sectors.get(_key(value))
    if result is None:
        raise ValueError(f"Sector laboral no reconocido: {value}")
    return result


def _company_nit(source: str, name: str) -> str:
    digest = hashlib.sha256(f"{source}:{name}".encode("utf-8")).hexdigest()[:20]
    return f"EXT-{digest}"


def template_csv() -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(IMPORT_HEADERS)
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def import_jobs(
    filename: str,
    content: bytes,
    current_user: User,
    db: UnitOfWork,
    *,
    nlp: TextAnalysis,
) -> dict:
    rows = _rows(filename, content)
    return _import_rows(rows, current_user, db, nlp=nlp)


def sync_job_catalog(
    *,
    keywords: str,
    location: str,
    pages: int,
    result_count: int,
    current_user: User,
    db: UnitOfWork,
    nlp: TextAnalysis,
    catalog: JobCatalog,
) -> dict:
    if not catalog.configured:
        raise UseCaseError(
            status_code=503,
            detail="Jooble aún no está conectado. Agrega JOOBLE_API_KEY en Azure.",
        )
    rows: list[dict[str, object]] = []
    provider_total = 0
    try:
        for page_number in range(1, pages + 1):
            result = catalog.search(
                keywords=keywords,
                location=location,
                page=page_number,
                result_count=result_count,
            )
            provider_total = result.total
            rows.extend(_catalog_row(job) for job in result.jobs)
            if not result.jobs or len(rows) >= result.total:
                break
    except CatalogProviderError as error:
        raise UseCaseError(status_code=502, detail=str(error)) from error

    if not rows:
        raise UseCaseError(
            status_code=404,
            detail="Jooble no encontró vacantes para esa búsqueda y ubicación.",
        )
    result = _import_rows(rows, current_user, db, nlp=nlp)
    return {
        **result,
        "fetched": len(rows),
        "provider_total": provider_total,
        "source": "Jooble",
    }


def _import_rows(
    rows: list[dict[str, object]],
    current_user: User,
    db: UnitOfWork,
    *,
    nlp: TextAnalysis,
) -> dict:
    if not rows:
        raise UseCaseError(status_code=400, detail="El archivo no contiene vacantes")
    if len(rows) > MAX_IMPORT_ROWS:
        raise UseCaseError(
            status_code=400,
            detail=f"Cada archivo puede contener máximo {MAX_IMPORT_ROWS} vacantes",
        )

    created = updated = companies_created = 0
    errors: list[dict[str, object]] = []
    seen_keys: set[tuple[str, str]] = set()
    now = datetime.utcnow()
    modality_map = {
        "remoto": "remote",
        "remote": "remote",
        "hibrido": "hybrid",
        "hybrid": "hybrid",
        "presencial": "onsite",
        "onsite": "onsite",
    }
    contract_map = {
        "tiempo_completo": "full_time",
        "full_time": "full_time",
        "medio_tiempo": "part_time",
        "part_time": "part_time",
        "contrato": "contract",
        "contract": "contract",
        "practicas": "internship",
        "practica": "internship",
        "internship": "internship",
        "temporal": "temporary",
        "temporary": "temporary",
    }

    for row_number, row in enumerate(rows, start=2):
        try:
            source = _text(row, "fuente", required=True)[:80]
            external_id = _text(row, "id_externo", required=True)[:180]
            identity = (source.casefold(), external_id.casefold())
            if identity in seen_keys:
                raise ValueError("La vacante está repetida dentro del archivo")
            seen_keys.add(identity)
            company_name = _text(row, "empresa", required=True)
            title = _text(row, "cargo", required=True)
            description = _text(row, "descripcion", required=True)
            requirements = _text(row, "requisitos", required=True)
            if len(company_name) > 180 or len(title) > 180:
                raise ValueError(
                    "El nombre de la empresa o el cargo supera 180 caracteres"
                )
            if len(description) < 10 or len(requirements) < 10:
                raise ValueError(
                    "La descripción y los requisitos necesitan al menos 10 caracteres"
                )
            sector = _sector(_text(row, "sector", required=True))
            external_url = _url(
                row.get("enlace_externo"), "enlace_externo", required=True
            )
            company_url = _url(row.get("sitio_empresa"), "sitio_empresa")
            published_at = _date(row.get("fecha_publicacion"), "fecha_publicacion")
            expires_at = _date(row.get("fecha_vencimiento"), "fecha_vencimiento")
            if expires_at and expires_at < now:
                raise ValueError("La fecha de vencimiento ya pasó")
            modality = _mapped(
                _text(row, "modalidad"), modality_map, "onsite", "modalidad"
            )
            employment_type = _mapped(
                _text(row, "tipo_contrato"),
                contract_map,
                "full_time",
                "tipo_contrato",
            )
            explicit_skills = _split(row.get("habilidades"))
            benefits = _split(row.get("beneficios"))
            languages = _languages(row.get("idiomas"))
            analysis = nlp.analyze_job(f"{title} {description} {requirements}")
            skills = list(dict.fromkeys([*explicit_skills, *analysis.skills]))[:40]

            company = db.companies.find_external(source, company_name)
            if company is None:
                company = db.companies.new(
                    owner_user_id=current_user.id,
                    name=company_name,
                    nit=_company_nit(source, company_name),
                    description=_text(row, "descripcion_empresa")
                    or f"Empresa incluida mediante el catálogo autorizado de {source}.",
                    website=company_url,
                    sector=sector,
                    location=", ".join(
                        filter(
                            None,
                            [_text(row, "ciudad"), _text(row, "departamento")],
                        )
                    )
                    or None,
                    is_external=True,
                    source_name=source,
                )
                db.add(company)
                db.flush()
                companies_created += 1
            else:
                if company_url:
                    company.website = company_url
                company.sector = company.sector or sector

            values = {
                "company_id": company.id,
                "title": title,
                "description": description,
                "requirements": requirements,
                "salary": _salary(row.get("salario")),
                "location": _text(row, "ciudad") or None,
                "department": _text(row, "departamento") or None,
                "modality": modality,
                "employment_type": employment_type,
                "sector": sector,
                "status": "active",
                "benefits": benefits,
                "languages": languages,
                "skills": skills,
                "embedding": analysis.embedding,
                "source_kind": "external",
                "source_name": source,
                "external_id": external_id,
                "external_url": external_url,
                "expires_at": expires_at,
                "last_seen_at": now,
            }
            job = db.jobs.find_external(source, external_id)
            if job is None:
                job = db.jobs.new(**values, created_at=published_at or now)
                db.add(job)
                created += 1
            else:
                for key, value in values.items():
                    setattr(job, key, value)
                if published_at:
                    job.created_at = published_at
                updated += 1
        except ValueError as error:
            errors.append({"row": row_number, "detail": str(error)})

    if not created and not updated:
        raise UseCaseError(
            status_code=422,
            detail={
                "message": "Ninguna vacante pudo importarse",
                "errors": errors[:50],
            },
        )
    db.commit()
    return {
        "created": created,
        "updated": updated,
        "companies_created": companies_created,
        "rejected": len(errors),
        "errors": errors[:50],
    }


def _catalog_row(job: CatalogJob) -> dict[str, object]:
    city, department = _location_parts(job.location)
    text = f"{job.title} {job.description}".casefold()
    return {
        "fuente": "Jooble",
        "id_externo": job.external_id,
        "empresa": job.company,
        "cargo": job.title,
        "descripcion": job.description
        or "Consulta la descripción completa de esta oportunidad en Jooble.",
        "requisitos": (
            "La API de Jooble entrega un resumen del anuncio. Confirma los requisitos "
            "completos y las condiciones en la publicación original."
        ),
        "sector": _infer_sector(text),
        "enlace_externo": job.url,
        "fecha_publicacion": job.published_at,
        "fecha_vencimiento": datetime.utcnow() + timedelta(days=30),
        "salario": job.salary,
        "departamento": department,
        "ciudad": city,
        "modalidad": _infer_modality(text),
        "tipo_contrato": _infer_contract(job.employment_type),
        "descripcion_empresa": (
            "Empresa con una oportunidad publicada en el catálogo autorizado de Jooble."
        ),
    }


def _location_parts(location: str) -> tuple[str, str]:
    parts = [part.strip() for part in location.split(",") if part.strip()]
    if len(parts) >= 2:
        return parts[0], parts[-1]
    return (parts[0], "") if parts else ("", "")


def _infer_modality(text: str) -> str:
    if any(term in text for term in ("remoto", "remote", "teletrabajo")):
        return "remoto"
    if any(term in text for term in ("híbrido", "hibrido", "hybrid")):
        return "híbrido"
    return "presencial"


def _infer_contract(value: str) -> str:
    normalized = _key(value)
    if any(term in normalized for term in ("part_time", "medio_tiempo")):
        return "medio tiempo"
    if any(term in normalized for term in ("intern", "practica", "pasantia")):
        return "prácticas"
    if any(term in normalized for term in ("temporary", "temporal")):
        return "temporal"
    if any(term in normalized for term in ("contract", "contrato")):
        return "contrato"
    return "tiempo completo"


def _infer_sector(text: str) -> str:
    rules = (
        ("Salud y bienestar", ("salud", "médic", "medic", "enfermer", "clínic")),
        ("Educacion", ("docente", "profesor", "educación", "colegio", "universidad")),
        ("Finanzas y banca", ("contable", "contador", "financ", "banco", "tesorer")),
        (
            "Ventas y comercio",
            (
                "venta",
                "comercial",
                "vendedor",
                "asesor comercial",
                "servicio al cliente",
            ),
        ),
        (
            "Logistica y transporte",
            ("logística", "logistica", "conductor", "transporte", "bodega"),
        ),
        ("Recursos humanos", ("recursos humanos", "talento humano", "reclut")),
        (
            "Marketing y publicidad",
            ("marketing", "mercadeo", "publicidad", "contenido"),
        ),
        ("Turismo y hoteleria", ("hotel", "turismo", "restaurante", "cocina")),
        (
            "Construccion e ingenieria",
            ("construcción", "construccion", "obra", "ingeniero civil"),
        ),
        ("Manufactura", ("producción", "produccion", "operario", "planta")),
        (
            "Tecnologia y software",
            (
                "software",
                "desarrollador",
                "programador",
                "sistemas",
                "datos",
                "soporte técnico",
            ),
        ),
        (
            "Sector publico y social",
            ("trabajo social", "fundación", "fundacion", "comunitario"),
        ),
    )
    for sector, terms in rules:
        if any(term in text for term in terms):
            return sector
    return "Retail y consumo masivo"
